"""
DailyKhata Word Document Export Generator
Converts D1.md, D2.md, and D3.md into polished Microsoft Word (.docx) files
following 00_Project_Governance/DOCUMENTATION_STANDARD.md.
"""

import re
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

REPO_ROOT = Path("d:/GitHub/project-part-I")
EXPORTS_DIR = REPO_ROOT / "exports"
ASSETS_LOGO = REPO_ROOT / "assets/bnu-logo/BNU Logo XXL Black smol@0.1x.png"

TEAM = [
    ("Saad Mughal", "F2023-009"),
    ("Khadijah Zahoor", "F2023-956"),
    ("Zainab Ali Baig", "F2023-528"),
    ("Aleeza Qaiser", "F2023-611"),
]
SUPERVISOR = "Nouman Ali"
ACADEMIC_YEAR = "2026-2027"
SUBMISSION_DATE = "September 2026"

DELIVERABLES_META = {
    1: {
        "src": REPO_ROOT / "01_Deliverable_1/D1.md",
        "out": EXPORTS_DIR / "DailyKhata_Deliverable_1.docx",
        "title": "Deliverable 1: Target Market & Audience Analysis",
        "version": "2.0",
        "doc_id": "D1",
    },
    2: {
        "src": REPO_ROOT / "02_Deliverable_2/D2.md",
        "out": EXPORTS_DIR / "DailyKhata_Deliverable_2.docx",
        "title": "Deliverable 2: Problem Identification, Stakeholder Identification & Requirement Elicitation",
        "version": "2.0",
        "doc_id": "D2",
    },
    3: {
        "src": REPO_ROOT / "03_Deliverable_3/D3.md",
        "out": EXPORTS_DIR / "DailyKhata_Deliverable_3.docx",
        "title": "Deliverable 3: Requirements Specification",
        "version": "1.0",
        "doc_id": "D3",
    },
}

def el(tag, **attrs):
    e = OxmlElement("w:" + tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e

def add_field(p, code):
    r = p.add_run()
    r._r.append(el("fldChar", fldCharType="begin"))
    r = p.add_run()
    e = el("instrText")
    e.set(qn("xml:space"), "preserve")
    e.text = " " + code + " "
    r._r.append(e)
    p.add_run()._r.append(el("fldChar", fldCharType="separate"))
    p.add_run("")
    p.add_run()._r.append(el("fldChar", fldCharType="end"))

def setup_page_and_styles(doc):
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.25)
    sec.right_margin = Inches(1.0)
    sec.header_distance = Inches(0.45)
    sec.footer_distance = Inches(0.45)
    sec.different_first_page_header_footer = True

    # Configure styles to Times New Roman, black
    for s in doc.styles:
        if s.type in (1, 2):
            s.font.name = "Times New Roman"
            s.font.color.rgb = RGBColor(0, 0, 0)
            rp = s.element.get_or_add_rPr()
            rf = rp.find(qn("w:rFonts"))
            if rf is not None:
                for key in list(rf.attrib):
                    if key.endswith("Theme"):
                        del rf.attrib[key]
                for key in ["ascii", "hAnsi", "eastAsia", "cs"]:
                    rf.set(qn("w:" + key), "Times New Roman")
            s.element.get_or_add_rPr().append(el("lang", val="en-GB"))

    # Normal body style
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.space_before = Pt(0)

    # Heading styles
    sizes = {
        "Heading 1": (16, Pt(12), Pt(6)),
        "Heading 2": (14, Pt(12), Pt(6)),
        "Heading 3": (12, Pt(12), Pt(6)),
        "Heading 4": (12, Pt(12), Pt(6)),
    }
    for name, (size, before, after) in sizes.items():
        if name in doc.styles:
            s = doc.styles[name]
            s.font.name = "Times New Roman"
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.italic = (name == "Heading 4")
            s.paragraph_format.space_before = before
            s.paragraph_format.space_after = after
            s.paragraph_format.keep_with_next = True
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Header and Footer
    hdr = sec.header.paragraphs[0]
    hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hdr.paragraph_format.space_after = Pt(0)
    hrun = hdr.add_run("DailyKhata | Final Year Project")
    hrun.font.name = "Times New Roman"
    hrun.font.size = Pt(9)
    hrun.font.color.rgb = RGBColor(128, 128, 128)

    ftr = sec.footer.paragraphs[0]
    ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ftr.paragraph_format.space_after = Pt(0)
    frun1 = ftr.add_run("Beaconhouse National University | BS Computer Science\n")
    frun1.font.name = "Times New Roman"
    frun1.font.size = Pt(9)
    frun2 = ftr.add_run("Page ")
    frun2.font.name = "Times New Roman"
    frun2.font.size = Pt(9)
    add_field(ftr, "PAGE")
    frun3 = ftr.add_run(" of ")
    frun3.font.name = "Times New Roman"
    frun3.font.size = Pt(9)
    add_field(ftr, "NUMPAGES")
    for r in ftr.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(9)

    doc.settings.element.append(el("updateFields", val="true"))

def add_cover_page(doc, title, subtitle_meta):
    # Top University Block
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_uni.paragraph_format.space_before = Pt(12)
    p_uni.paragraph_format.space_after = Pt(4)
    r = p_uni.add_run("BEACONHOUSE NATIONAL UNIVERSITY\n")
    r.bold = True
    r.font.size = Pt(14)
    r2 = p_uni.add_run("School of Computer Science | Final Year Project")
    r2.font.size = Pt(11)

    # Logo
    if ASSETS_LOGO.exists():
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(24)
        p_logo.paragraph_format.space_after = Pt(24)
        run_logo = p_logo.add_run()
        run_logo.add_picture(str(ASSETS_LOGO), width=Inches(1.27), height=Inches(1.5))

    # Project Title
    p_proj = doc.add_paragraph()
    p_proj.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_proj.paragraph_format.space_before = Pt(12)
    p_proj.paragraph_format.space_after = Pt(6)
    r_proj = p_proj.add_run("DAILYKHATA")
    r_proj.bold = True
    r_proj.font.size = Pt(22)

    # Deliverable Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(6)
    p_sub.paragraph_format.space_after = Pt(36)
    r_sub = p_sub.add_run(subtitle_meta)
    r_sub.bold = True
    r_sub.font.size = Pt(14)

    # Team Members Block
    p_team_hdr = doc.add_paragraph()
    p_team_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_team_hdr.paragraph_format.space_before = Pt(18)
    p_team_hdr.paragraph_format.space_after = Pt(4)
    r_th = p_team_hdr.add_run("Project Team:")
    r_th.bold = True
    r_th.font.size = Pt(11)

    p_team = doc.add_paragraph()
    p_team.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_team.paragraph_format.space_before = Pt(0)
    p_team.paragraph_format.space_after = Pt(24)
    for i, (name, sid) in enumerate(TEAM):
        p_team.add_run(f"{name} ({sid})" + ("\n" if i < len(TEAM)-1 else "")).font.size = Pt(11)

    # Academic metadata
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(12)
    p_meta.paragraph_format.space_after = Pt(0)
    p_meta.add_run(f"Supervisor: {SUPERVISOR}\nAcademic Year: {ACADEMIC_YEAR}\nDate: {SUBMISSION_DATE}").font.size = Pt(10.5)

    doc.add_page_break()

def add_document_control(doc, title, version):
    p_hdr = doc.add_paragraph()
    p_hdr.paragraph_format.space_before = Pt(12)
    p_hdr.paragraph_format.space_after = Pt(12)
    r = p_hdr.add_run("Document Control")
    r.bold = True
    r.font.size = Pt(16)

    controls = [
        ("Project", "DailyKhata"),
        ("Document Title", title),
        ("Version", version),
        ("Document Status", "Official University Submission Edition"),
        ("Authors", "DailyKhata Project Team"),
        ("Supervisor", SUPERVISOR),
        ("Academic Year", ACADEMIC_YEAR),
        ("Submission Date", SUBMISSION_DATE),
    ]

    t = doc.add_table(rows=len(controls), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    col_widths = [Inches(1.8), Inches(4.22)]

    t._tbl.tblPr.append(el("tblLayout", type="fixed"))
    tblGrid = el("tblGrid")
    for w in col_widths:
        tblGrid.append(el("gridCol", w=str(int(w.inches * 1440))))
    t._tbl.insert(1, tblGrid)

    for row_idx, (k, v) in enumerate(controls):
        row = t.rows[row_idx]
        row._tr.get_or_add_trPr().append(el("cantSplit"))
        cell_k, cell_v = row.cells[0], row.cells[1]
        cell_k.width = col_widths[0]
        cell_v.width = col_widths[1]

        p_k = cell_k.paragraphs[0]
        p_k.paragraph_format.space_before = Pt(3)
        p_k.paragraph_format.space_after = Pt(3)
        p_k.paragraph_format.line_spacing = 1.05
        rk = p_k.add_run(k)
        rk.bold = True
        rk.font.name = "Times New Roman"
        rk.font.size = Pt(10.5)

        p_v = cell_v.paragraphs[0]
        p_v.paragraph_format.space_before = Pt(3)
        p_v.paragraph_format.space_after = Pt(3)
        p_v.paragraph_format.line_spacing = 1.05
        rv = p_v.add_run(v)
        rv.font.name = "Times New Roman"
        rv.font.size = Pt(10.5)

        # Style borders and margins
        for c, w in ((cell_k, col_widths[0]), (cell_v, col_widths[1])):
            tcPr = c._tc.get_or_add_tcPr()
            tcPr.append(el("tcW", w=str(int(w.inches * 1440)), type="dxa"))
            mar = el("tcMar")
            for s in ["top", "bottom"]:
                mar.append(el(s, w=85, type="dxa"))
            for s in ["left", "right"]:
                mar.append(el(s, w=120, type="dxa"))
            tcPr.append(mar)

    # Table borders
    borders = el("tblBorders")
    for s in ["top", "left", "bottom", "right", "insideH"]:
        borders.append(el(s, val="single", sz=4, color="D9D9D9"))
    borders.append(el("insideV", val="single", sz=4, color="D9D9D9"))
    t._tbl.tblPr.append(borders)

    doc.add_page_break()

def add_table_of_contents(doc):
    p_hdr = doc.add_paragraph()
    p_hdr.paragraph_format.space_before = Pt(12)
    p_hdr.paragraph_format.space_after = Pt(12)
    r = p_hdr.add_run("Table of Contents")
    r.bold = True
    r.font.size = Pt(16)

    p_toc = doc.add_paragraph()
    p_toc.paragraph_format.space_before = Pt(0)
    p_toc.paragraph_format.space_after = Pt(12)
    add_field(p_toc, 'TOC \\o "1-3" \\h \\z \\u')

    doc.add_page_break()

def format_inline_text(paragraph, text):
    """Parses **bold**, *italic*, and regular text runs into paragraph."""
    # Pattern to match bold (**text**) or italic (*text*)
    tokens = re.split(r'(\*\*[^*]+?\*\*|\*[^*]+?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            r = paragraph.add_run(token[2:-2])
            r.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.italic = True
        else:
            paragraph.add_run(token)

def get_table_column_widths(headers):
    """
    Returns explicit column widths totaling 6.02 inches for all known tables
    across D1, D2, and D3, with a robust fallback.
    """
    norm_headers = [h.strip().lower() for h in headers]
    hdr_tuple = tuple(norm_headers)

    # 1. D1 tables
    if hdr_tuple == ("persona", "age and role", "goals, behaviour and expectations"):
        return [Inches(1.5), Inches(1.8), Inches(2.72)]
    if hdr_tuple == ("user need", "expected response from dailykhata"):
        return [Inches(2.0), Inches(4.02)]
    if hdr_tuple == ("priority", "capabilities"):
        return [Inches(1.7), Inches(4.32)]

    # 2. D2 tables
    if hdr_tuple == ("existing method", "limitation in the reported workflow", "implication for dailykhata"):
        return [Inches(1.4), Inches(2.31), Inches(2.31)]
    if hdr_tuple == ("evidence basis", "main finding", "requirement implication"):
        return [Inches(1.4), Inches(2.31), Inches(2.31)]
    if hdr_tuple == ("stage", "dailykhata application"):
        return [Inches(1.5), Inches(4.52)]
    if hdr_tuple == ("stakeholder and role", "interest and needs", "impact on the proposed system"):
        return [Inches(1.5), Inches(2.26), Inches(2.26)]
    if hdr_tuple == ("candidate", "basis and proposed behaviour", "priority"):
        return [Inches(1.5), Inches(3.52), Inches(1.0)]
    if hdr_tuple == ("quality area", "expected behaviour", "evaluation needed"):
        return [Inches(1.4), Inches(2.31), Inches(2.31)]

    # 3. D3 tables
    if len(headers) == 3 and norm_headers[0] == "id" and "priority" in norm_headers[2]:
        return [Inches(1.15), Inches(3.87), Inches(1.0)]
    if hdr_tuple == ("priority", "business requirements", "functional requirements", "non-functional requirements"):
        return [Inches(1.52), Inches(1.5), Inches(1.5), Inches(1.5)]
    if hdr_tuple == ("d2 requirement candidate", "business requirements", "functional requirements"):
        return [Inches(2.02), Inches(1.8), Inches(2.2)]

    # General Fallback
    usable_width = 6.02
    n = len(headers)
    if n == 2:
        return [Inches(1.8), Inches(4.22)]
    elif n == 3:
        return [Inches(1.5), Inches(2.5), Inches(2.02)]
    elif n == 4:
        return [Inches(1.5), Inches(1.5), Inches(1.5), Inches(1.52)]
    else:
        w = usable_width / max(n, 1)
        return [Inches(w)] * n

def add_clean_table(doc, table_data):
    """
    Renders markdown table as Word table with light grey shading and borders,
    repeating header, cantSplit rows, and balanced column widths.
    """
    if not table_data or len(table_data) < 2:
        return

    headers = table_data[0]
    rows = table_data[1:]
    col_count = len(headers)

    t = doc.add_table(rows=1 + len(rows), cols=col_count)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    col_widths = get_table_column_widths(headers)

    # Set fixed table layout
    t._tbl.tblPr.append(el("tblLayout", type="fixed"))

    # Add tblGrid with explicit gridCol widths
    tblGrid = el("tblGrid")
    for w in col_widths:
        tblGrid.append(el("gridCol", w=str(int(w.inches * 1440))))
    t._tbl.insert(1, tblGrid)

    # Style header row
    hdr_row = t.rows[0]
    hdr_row._tr.get_or_add_trPr().append(el("tblHeader"))
    hdr_row._tr.get_or_add_trPr().append(el("cantSplit"))

    for i, h_text in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.width = col_widths[i]
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(el("tcW", w=str(int(col_widths[i].inches * 1440)), type="dxa"))
        tcPr.append(el("shd", fill="E7E7E7"))
        mar = el("tcMar")
        for s in ["top", "bottom"]:
            mar.append(el(s, w=90, type="dxa"))
        for s in ["left", "right"]:
            mar.append(el(s, w=120, type="dxa"))
        tcPr.append(mar)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05
        format_inline_text(p, h_text.strip())
        for run in p.runs:
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)

    # Style data rows
    for r_idx, row_data in enumerate(rows):
        row = t.rows[1 + r_idx]
        row._tr.get_or_add_trPr().append(el("cantSplit"))
        for i in range(col_count):
            val = row_data[i] if i < len(row_data) else ""
            cell = row.cells[i]
            cell.width = col_widths[i]
            tcPr = cell._tc.get_or_add_tcPr()
            tcPr.append(el("tcW", w=str(int(col_widths[i].inches * 1440)), type="dxa"))
            mar = el("tcMar")
            for s in ["top", "bottom"]:
                mar.append(el(s, w=85, type="dxa"))
            for s in ["left", "right"]:
                mar.append(el(s, w=120, type="dxa"))
            tcPr.append(mar)

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.05
            format_inline_text(p, val.strip())
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)

    # Add light borders
    borders = el("tblBorders")
    for s in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        borders.append(el(s, val="single", sz=4, color="D9D9D9"))
    t._tbl.tblPr.append(borders)

    # Spacing after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def parse_markdown_and_build(doc, md_text):
    lines = md_text.splitlines()
    in_table = False
    table_data = []

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle Table parsing
        if stripped.startswith("|") and "|" in stripped[1:]:
            in_table = True
            # Split cells by '|', ignore first and last empty
            parts = [c.strip() for c in stripped.split("|")[1:-1]]
            # If it's a separator line like | --- | --- |, skip
            if not all(re.match(r'^:?-+:?$', p) for p in parts):
                table_data.append(parts)
            i += 1
            continue
        else:
            if in_table:
                # Table ended, render it
                add_clean_table(doc, table_data)
                table_data = []
                in_table = False

        if not stripped:
            i += 1
            continue

        # Skip the top document h1 (already handled in cover/doc control)
        if stripped.startswith("# ") and not stripped.startswith("## "):
            i += 1
            continue

        # Headings
        if stripped.startswith("#### "):
            h_text = stripped[5:].strip()
            p = doc.add_heading(h_text, level=3)
            p.paragraph_format.keep_with_next = True
            i += 1
            continue
        elif stripped.startswith("### "):
            h_text = stripped[4:].strip()
            p = doc.add_heading(h_text, level=2)
            p.paragraph_format.keep_with_next = True
            i += 1
            continue
        elif stripped.startswith("## "):
            h_text = stripped[3:].strip()
            p = doc.add_heading(h_text, level=1)
            p.paragraph_format.keep_with_next = True
            i += 1
            continue

        # Blockquote
        if stripped.startswith("> "):
            bq_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            format_inline_text(p, bq_text)
            for r in p.runs:
                r.italic = True
                r.font.size = Pt(11)
            i += 1
            continue

        # Bullet List
        if stripped.startswith("- ") or stripped.startswith("* "):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            format_inline_text(p, bullet_text)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
            i += 1
            continue

        # Numbered List (e.g. 1. text)
        m_num = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m_num:
            num_idx, num_text = m_num.groups()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            r_num = p.add_run(f"{num_idx}. ")
            r_num.bold = True
            format_inline_text(p, num_text)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
            i += 1
            continue

        # Normal Paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        format_inline_text(p, stripped)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)

        i += 1

    if in_table:
        add_clean_table(doc, table_data)

def generate_deliverable(deliv_id):
    meta = DELIVERABLES_META[deliv_id]
    src_path = meta["src"]
    out_path = meta["out"]

    print(f"Generating {meta['doc_id']}: {src_path.name} -> {out_path.name}...")
    md_text = src_path.read_text(encoding="utf-8")

    doc = Document()
    setup_page_and_styles(doc)

    # 1. Cover Page
    add_cover_page(doc, "DailyKhata", meta["title"])

    # 2. Document Control
    add_document_control(doc, meta["title"], meta["version"])

    # 3. Table of Contents
    add_table_of_contents(doc)

    # 4. Main Body Content
    parse_markdown_and_build(doc, md_text)

    # Ensure output dir exists
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"Saved {out_path.name} successfully.")

def main():
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    for d_id in [1, 2, 3]:
        generate_deliverable(d_id)
    print("\nAll 3 deliverables generated successfully!")

if __name__ == "__main__":
    main()
