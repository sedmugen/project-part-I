from pathlib import Path
from docx import Document
from pypdf import PdfReader
import json
w=Path('D:/BNU/SM7/FYP/03_Deliverable_3/requirements_work')
p=w.parent/'DailyKhata_Deliverable_3_v1.0.docx'
(w/'before_polish_pages.json').write_text(json.dumps([x.extract_text() for x in PdfReader(w/'D3_QA.pdf').pages]),encoding='utf-8')
d=Document(p)
for t in d.tables:
 if t.rows[0].cells[0].text not in ('Scenario','Issue'):continue
 for row in t.rows[1:]:
  cell=row.cells[0];value=cell.text.split()[0];runs=[r for p in cell.paragraphs for r in p.runs]
  for r in runs:r.text=''
  runs[0].text=value
for para in d.paragraphs:
 if para.text.startswith('Scope and downstream behaviour: SC-02, SC-05, SC-13, SC-24.'):
  for run in para.runs:run.text=run.text.replace('SC-13, SC-24.','SC-13, SC-24, SC-30.')
d.save(p)
print('Two narrow label columns simplified; BR-006 direct Android scope link added.')
