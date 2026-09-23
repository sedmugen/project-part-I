from pathlib import Path
import json, re, sys
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('D:/BNU/SM7/FYP')
WORK=ROOT/'00_Project_Governance/reconstruction_work'
DATE='23 September 2026'
TEAM=['Saad Mughal (F2023-009)','Khadijah Zahoor (F2023-956)','Zainab Ali Baig (F2023-528)','Aleeza Qaiser (F2023-611)']
DEFINITION='DailyKhata is a personal expense management application primarily designed for Pakistani users who frequently deal with paper receipts, cash transactions and informal expense tracking. It reduces the effort required to maintain personal or household expense records through receipt digitisation, OCR, quick cash expense capture, manual expense entry, automatic categorisation, user confirmation and correction, a structured Chart of Accounts, transaction history, spending summaries, item-level price history, semantic product classification and an AI Financial Assistant.'
PROBLEM='Paper receipts and unrecorded cash purchases leave personal and household expense histories incomplete. Manual upkeep makes consistent recording difficult. Users need low-effort capture and reliable ways to review and understand spending.'
SCOPE=[
('SC-01','User registration and authentication','Must Have'),('SC-02','Basic user profile and settings','Must Have'),('SC-03','Personal digital expense ledger','Must Have'),('SC-04','Photograph or upload receipts','Must Have'),('SC-05','OCR processing','Must Have'),('SC-06','Structured extraction of merchant and date where available, line items, quantities where available, prices and receipt total','Must Have'),('SC-07','Review OCR results and edit incorrect extracted values before saving','Must Have'),('SC-08','Manual quick expense entry','Must Have'),('SC-09','Text-based cash expense entry','Must Have'),('SC-10','Voice-based cash expense entry','Should Have'),('SC-11','Automatic transaction and line-item categorisation','Must Have'),('SC-12','User confirmation or correction of AI-generated categories','Must Have'),('SC-13','Locally understandable expense category labels','Must Have'),('SC-14','Full internal Chart of Accounts','Must Have'),('SC-15','Mapping of user-facing categories to accounting accounts','Must Have'),('SC-16','Accounting-compatible transaction storage','Must Have'),('SC-17','Transaction history','Must Have'),('SC-18','Search and filtering','Must Have'),('SC-19','Category-level spending totals','Must Have'),('SC-20','Basic monthly spending summaries','Must Have'),('SC-21','Item-level price history','Must Have'),('SC-22','Semantic product classification','Must Have'),('SC-23','Natural-language financial querying','Must Have'),('SC-24','AI Financial Assistant','Must Have'),('SC-25','Dashboard','Must Have'),('SC-26','Editing and deleting saved expense transactions','Must Have'),('SC-27','Secure user-data storage','Must Have'),('SC-28','English and Urdu localisation','Should Have'),('SC-29','CSV and/or PDF export','Should Have'),('SC-30','Reasonable support for ordinary Android-class devices','Must Have'),('SS-03','Custom category labels','Should Have'),('SS-04','Simple spending visualisations','Should Have'),('SS-05','Month-to-month comparison views','Should Have'),('SS-06','Receipt image archive','Should Have'),('SS-08','User-created categories','Could Have'),('SS-09','Offline draft expense entry','Could Have'),('SS-10','Expense reminders','Could Have'),('SS-11','Exportable household summaries','Could Have')]
EXCLUDED='Direct bank integration; automatic debit-card or credit-card synchronisation; Easypaisa and JazzCash transaction integration; open banking; payment processing; money transfers; inventory management; point-of-sale functionality; payroll; enterprise ERP; complete SME accounting; tax filing; investment management; credit scoring; fraud-detection infrastructure; commercial banking functionality; enterprise multi-tenant organisation management; banking-grade regulatory certification; production infrastructure for millions of concurrent users; arbitrary third-party financial integrations.'
FUTURE='Bank transaction importing; wallet integrations; shared household accounts; collaborative budgeting; business accounting features; inventory; POS; recurring expenses; budgeting goals; predictive analytics; anomaly detection; advanced financial recommendations; deeper Urdu support; additional languages; multi-currency support; retailer integrations; e-receipt ingestion.'
PERSONAS=[
('P-01','Ayesha Tariq','34','Homemaker and household budget manager','Primary','Track weekly household spending and grocery prices','Handles paper receipts and cash; prefers a photograph to transcription','Needs readable review, familiar categories and price history; expects correction and privacy'),
('P-02','Bilal Ahmed','27','Salaried professional','Primary','Reconcile and understand monthly personal expenses','Comfortable with digital tools; reviews spending periodically','Needs searchable history and plain-language answers; expects minimal upkeep'),
('P-03','Hamza Sheikh','21','University student','Primary','Track personal spending and allowance with little effort','Frequent small cash purchases; low tolerance for repeated form entry','Needs quick text entry without a receipt; expects useful category totals'),
('P-04','Zubair Hassan','Not recorded','Family member and spouse of a primary user','Secondary beneficiary','Review household spending with the account holder','Occasional review of summaries shared by the account holder','Needs understandable totals; no separate household login or permissions implied'),
('P-05','Usman Raza','41','Small shopkeeper','Secondary direct user','Track petty cash and supplier expenses','Uses the same receipt and expense workflow for a general store','Needs dependable expense records; sales, inventory and POS are excluded')]
DECISIONS=[
('DEC-01','D1 persona table gives Usman age 41; D2 pp. 3, 7, 8 and 10 give 46.','Retain 41 as the canonical design persona age; preserve the source discrepancy in the evidence register.','Earlier D1 fact takes precedence without corroborating participant records. Actual participant age needs team confirmation.','D1, D2, terminology, audit'),
('DEC-02','D1 summary narrows primary users to two groups; its main text includes salaried professionals. Hamza appears only in D2.','Three primary groups: homemakers and household budget managers, university students, salaried professionals. Carry Hamza Sheikh, 21, into D1.','Approved scope and D1 main text establish all three groups; D2 provides the student persona.','D1, D2, scope, terminology'),
('DEC-03','D1 calls Zubair a spouse but uses an inconsistent gendered description; both sources imply a shared budget.','Use gender-neutral family-member wording; review is with the account holder or through shared output. Age remains not recorded.','No evidence supports a new identity, age or multi-user household architecture.','D1, D2, scope, terminology'),
('DEC-04','D1 and D2 refer broadly to SMEs, daily transactions and business expenses.','Usman is a secondary direct user for petty cash and supplier expenses only.','Approved scope excludes complete SME accounting, sales, inventory and POS.','D1, D2, scope'),
('DEC-05','D2 p. 3 attributes a statement about a customer paying to Usman.','Preserve the statement as source-reported evidence, but do not derive sales or payment requirements.','The statement concerns cash capture friction; its income context falls outside expense scope.','D2, decision log, audit'),
('DEC-06','D1 promises categories without intervention; D2 requires confirmation.','Automatic categorisation proposes values; users confirm or correct before saving.','Approved scope and D2 elicitation explicitly require user control.','D1, D2, scope'),
('DEC-07','Voice/photo-first language and text-or-voice requirements blur mandatory and optional capture.','Manual quick entry and text-based cash entry are Must Have; voice-based cash entry is Should Have.','Preserves capture for no-receipt purchases without making speech recognition a baseline dependency.','D1, D2, scope'),
('DEC-08','D1 says several million users, typical Rs. 40,000 spending, universal pain and widespread abandonment within days or weeks.','Remove numerical market estimates and population-wide behavioural claims. Retain individual reported experiences with attribution.','No sampling basis, citation or market-size calculation exists in the workspace.','D1, D2, audit'),
('DEC-09','D1 and D2 make sweeping claims about Mint, Wallet, Cashew and bank apps.','Retain D2 comparison as a historical team-reported test only; do not state current competitor capabilities or superiority.','Versions, test dates, bank-app identity and screenshots are absent.','D1, D2, audit'),
('DEC-10','D1 says almost instant answers; D2 promises all capture under 10 seconds and OCR in a few seconds.','Treat speed as a usability expectation; measure capture, correction and service latency separately during later validation.','No benchmark supports a universal threshold; receipt complexity, network and corrections vary.','D1, D2, scope'),
('DEC-11','D2 expects operation with inconsistent connectivity; offline functionality is not clearly bounded.','Must Have workflows require honest loading, failure and retry behaviour; offline draft entry remains Could Have.','Ordinary-device support does not imply offline OCR or an offline language model.','D1, D2, scope'),
('DEC-12','Original deliverables omit or underdescribe authentication, ledger maintenance, full Chart of Accounts, semantic products and assistant breadth.','Carry the approved scope register in both deliverables and explain these capabilities as scope decisions.','The user-authorised scope and existing scope baseline supply their authority, not invented interviews.','All replacement documents'),
('DEC-13','Optional comparisons overlap with core assistant comparisons; future anomaly detection overlaps with unusual-spending discussion.','Assistant answers about periods, patterns and unusually high recorded spending are core. Dedicated comparison views are Should Have; predictive or automated anomaly-detection systems remain future.','Separates data-grounded conversational analysis from additional screens and predictive infrastructure.','D1, D2, scope, terminology'),
('DEC-14','Local labels, custom labels and user-created categories are conflated; SS duplicates SC entries.','Fixed local labels are Must Have; custom labels Should Have; new categories Could Have. SS-01 aliases SC-10, SS-02 aliases SC-29, SS-07 aliases SC-28.','Retains existing identifiers and avoids double-counting or priority drift.','D1, D2, scope, terminology'),
('DEC-15','D2 cover repeats the D1 title; original documents have inconsistent formatting and manual headings.','Use the D2 assignment title, one Word standard, automatic headings, contents and page fields. D1 title retains its original official wording.','Aligns covers with assignment purpose and supports future maintenance.','All replacement documents'),
('DEC-16','Existing governance is labelled Approved Baseline but no approval record is present. Submission dates are not established.','Preserve v1.0 files. Replacement baseline follows the user-authorised scope; supervisor endorsement is not claimed. Date replacements 23 September 2026.','An assignment deadline and file modification date do not establish actual submission or approval.','All replacement documents'),
('DEC-17','Standards-aligned accounting could imply certification or a complete business accounting suite.','Use stable accounts, account types, category mappings and accounting-compatible expense records. Balanced journals are mandatory wherever double-entry is used.','Supports structured financial records without claiming IFRS, banking or regulatory certification.','D1, D2, scope'),
('DEC-18','Interview and observation narratives lack underlying records; persona traits can be mistaken for verified participant facts.','Label evidence as source-reported; label personas as design profiles. Add no interviews, quotations, survey results or statistics.','Preserves provenance and avoids converting plausible analysis into fieldwork.','D1, D2, audit'),
('DEC-19','Existing standard prefers Roman front matter but permits continuous Arabic pagination.','Use continuous Arabic Page X of Y fields, counting the cover but hiding its number.','Permitted convention reduces numbering complexity and is identical across the new documents.','Documentation standard and all replacements'),
('DEC-20','Receipt retention is described as permanent; item normalisation could erase brand identity.','Structured records are core; image archive is Should Have. Preserve original receipt text separately from corrected and inferred product fields.','Storage lifetime is not guaranteed; semantic grouping must not merge distinct products or pack sizes.','D1, D2, scope, terminology')]

def el(tag,**attrs):
 e=OxmlElement('w:'+tag)
 for k,v in attrs.items(): e.set(qn('w:'+k),str(v))
 return e
def field(p,code):
 r=p.add_run(); r._r.append(el('fldChar',fldCharType='begin'))
 r=p.add_run(); e=el('instrText'); e.set(qn('xml:space'),'preserve'); e.text=' '+code+' '; r._r.append(e)
 p.add_run()._r.append(el('fldChar',fldCharType='separate'))
 p.add_run('')
 p.add_run()._r.append(el('fldChar',fldCharType='end'))

class Report:
 def __init__(self,title,filename,folder,deliverable=None,version='2.0'):
  self.d=Document(); self.title=title; self.path=ROOT/folder/filename; self.tc=0
  d=self.d; sec=d.sections[0]; sec.page_width=Mm(210); sec.page_height=Mm(297)
  sec.top_margin=Inches(1); sec.bottom_margin=Inches(1); sec.left_margin=Inches(1.25); sec.right_margin=Inches(1)
  sec.header_distance=Inches(.45); sec.footer_distance=Inches(.45); sec.different_first_page_header_footer=True
  for s in d.styles:
   if s.type in (1,2):
    s.font.name='Times New Roman'; s.font.color.rgb=RGBColor(0,0,0)
    rp=s.element.get_or_add_rPr()
    rf=rp.find(qn('w:rFonts'))
    if rf is not None:
     for key in list(rf.attrib):
      if key.endswith('Theme'): del rf.attrib[key]
     for key in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+key),'Times New Roman')
    for border in s.element.findall('.//'+qn('w:pBdr')): border.getparent().remove(border)
    s.element.get_or_add_rPr().append(el('lang',val='en-GB'))
  n=d.styles['Normal']; n.font.size=Pt(12); n.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; n.paragraph_format.line_spacing=1.15; n.paragraph_format.space_after=Pt(6)
  for i,size in enumerate([16,14,12,12],1):
   s=d.styles[f'Heading {i}']; s.font.size=Pt(size); s.font.bold=True; s.font.italic=(i==4)
   s.paragraph_format.space_before=Pt(12); s.paragraph_format.space_after=Pt(6); s.paragraph_format.keep_with_next=True
   s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
  for sn in ['Title','Subtitle']:
   d.styles[sn].font.size=Pt(20 if sn=='Title' else 16); d.styles[sn].font.bold=True; d.styles[sn].font.italic=False
  d.styles['Caption'].font.size=Pt(10); d.styles['Caption'].font.italic=False
  numbering=d.part.numbering_part.element
  ab=el('abstractNum',abstractNumId=50); ab.append(el('multiLevelType',val='multilevel'))
  for i in range(4):
   lvl=el('lvl',ilvl=i); lvl.append(el('start',val=1)); lvl.append(el('numFmt',val='decimal')); lvl.append(el('pStyle',val=f'Heading{i+1}')); lvl.append(el('lvlText',val='.'.join('%'+str(x+1) for x in range(i+1)))); lvl.append(el('suff',val='space')); ab.append(lvl)
   np=el('numPr'); np.append(el('ilvl',val=i)); np.append(el('numId',val=50)); d.styles[f'Heading {i+1}'].element.get_or_add_pPr().append(np)
  numbering.append(ab); num=el('num',numId=50); num.append(el('abstractNumId',val=50)); numbering.append(num)
  d.settings.element.append(el('updateFields',val='true'))
  p=sec.header.paragraphs[0]; p.text='DailyKhata | Final Year Project'; p.runs[0].font.size=Pt(9)
  p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Beaconhouse National University | BS Computer Science\n').font.size=Pt(9); p.add_run('Page ').font.size=Pt(9); field(p,'PAGE'); p.add_run(' of ').font.size=Pt(9); field(p,'NUMPAGES')
  for text,sty in [('Beaconhouse National University','Normal'),('BS Computer Science','Normal'),('Final Year Project','Normal'),('DAILYKHATA','Title'),((f'Deliverable {deliverable}\n' if deliverable else 'Project Governance\n')+title,'Subtitle')]:
   p=d.add_paragraph(text,sty); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(14)
  p=d.add_paragraph('Team members\n'+'\n'.join(TEAM)); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(18)
  p=d.add_paragraph('Supervisor: Nouman Ali\nAcademic Year: 2026-2027\nVersion '+version+' | '+DATE); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  d.add_page_break()
  self.front('Document Control'); self.p('Project: DailyKhata\nDocument: '+title+'\nVersion: '+version+'\nStatus: Authoritative reconstruction baseline under the user-authorised scope\nAuthors: DailyKhata Project Team\nSupervisor: Nouman Ali\nAcademic year: 2026-2027\nRevision date: '+DATE+'\nSubmission date: Not recorded for this replacement\nSupervisor endorsement: To be recorded after review')
  self.front('Revision History'); self.p('Historical source versions remain unchanged in their original locations. This replacement consolidates scope, terminology and evidence provenance under the recorded decisions. Version '+version+', '+DATE+', prepared for the DailyKhata Project Team.')
  d.add_page_break(); self.front('Table of Contents'); field(d.add_paragraph(),'TOC \\o "1-3" \\h \\z \\u'); d.add_page_break()
 def front(self,t):
  p=self.d.add_paragraph(t); p.paragraph_format.space_after=Pt(12); p.runs[0].bold=True; p.runs[0].font.size=Pt(16)
 def h(self,t,level=1): self.d.add_heading(t,level)
 def p(self,t):
  p=self.d.add_paragraph(t)
  if '\n' in t: p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 def table(self,caption,headers,rows,widths=None):
  self.tc+=1; p=self.d.add_paragraph(style='Caption'); p.paragraph_format.keep_with_next=True; p.add_run('Table '); field(p,'SEQ Table \\* ARABIC'); p.add_run(': '+caption)
  t=self.d.add_table(rows=1,cols=len(headers)); t.autofit=False
  if widths is None: widths=[6.02/len(headers)]*len(headers)
  for c,w in zip(t.columns,widths): c.width=Inches(w)
  for c,h,w in zip(t.rows[0].cells,headers,widths): c.text=h; c.width=Inches(w)
  t.rows[0]._tr.get_or_add_trPr().append(el('tblHeader'))
  for row in rows:
   for c,v,w in zip(t.add_row().cells,row,widths): c.text=str(v); c.width=Inches(w)
  for ri,row in enumerate(t.rows):
   row._tr.get_or_add_trPr().append(el('cantSplit'))
   for c in row.cells:
    pr=c._tc.get_or_add_tcPr(); margins=el('tcMar')
    for side in ['top','left','bottom','right']: margins.append(el(side,w=85,type='dxa'))
    pr.append(margins)
    if ri==0: pr.append(el('shd',fill='E7E7E7'))
    for p in c.paragraphs:
     p.alignment=WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.05
     if ri==0: p.paragraph_format.keep_with_next=True
     for r in p.runs: r.font.size=Pt(10); r.bold=ri==0
  borders=el('tblBorders')
  for side in ['top','left','bottom','right','insideH','insideV']: borders.append(el(side,val='single',sz=4,color='D9D9D9'))
  t._tbl.tblPr.append(borders)
  self.d.add_paragraph().paragraph_format.space_after=Pt(0)
 def finish(self):
  # A true Word list of tables, refreshed along with the TOC.
  if self.tc:
   idx=next(i for i,p in enumerate(self.d.paragraphs) if p.text=='Table of Contents')
   toc=self.d.paragraphs[idx+1]
   p=OxmlElement('w:p'); toc._p.addnext(p)
   from docx.text.paragraph import Paragraph
   q=Paragraph(p,toc._parent); q.add_run('List of Tables').bold=True; q.paragraph_format.space_before=Pt(14)
   p2=OxmlElement('w:p'); p.addnext(p2); field(Paragraph(p2,toc._parent),'TOC \\c "Table" \\h \\z')
  self.d.core_properties.title=self.title; self.d.core_properties.author='DailyKhata Project Team'; self.d.core_properties.subject='Authoritative DailyKhata project foundation'
  self.path.parent.mkdir(exist_ok=True); self.d.save(self.path)
  print(self.path)

def scope_table(r):
 r.p('The following register is shared by D1, D2 and Project Scope. Must Have means required for the minimum successful FYP; Should Have means intended but deferrable; Could Have means optional if time permits. Priority is not a claim of implementation.')
 r.table('Canonical feature scope and priorities',['ID','Capability','Priority'],SCOPE,[.66,4.41,.95])
 r.p('Legacy aliases are retained for traceability: SS-01 refers to SC-10, SS-02 to SC-29, and SS-07 to SC-28. They do not create additional features. Month-to-month comparison views are optional; answering period comparisons through the AI Financial Assistant is core.')

def persona_section(r,expanded=True):
 r.p('These are canonical design personas, not a new participant sample. Names and source-supported attributes are preserved from D1 and D2. Usman Raza uses the D1 design age of 41; D2 reports an interview age of 46 without corroborating records. DEC-01 records that unresolved source fact without leaving two competing persona definitions.')
 r.table('Canonical personas',['ID and name','Age and role','Classification'],[(p[0]+'\n'+p[1],p[2]+'\n'+p[3],p[4]) for p in PERSONAS],[1.55,3.12,1.35])
 if expanded:
  for p in PERSONAS:
   r.h(p[1],2); r.p('Goal: '+p[5]+'. Behaviour: '+p[6]+'. '+p[7]+'.')
  r.p('Digital confidence varies within every user group. Persona descriptions guide design and recruitment; they do not justify assuming that all homemakers have low digital literacy or all students have high technical skill. Bilal and Zubair are not reported interview participants. Zubair\'s age is not recorded and no age is inferred.')

def architecture(r):
 r.h('Accounting and category boundaries',2)
 r.p('DailyKhata must maintain a full internal Chart of Accounts with stable account identifiers, account types and mappings from user-facing expense categories. Full means a coherent internal account framework, including the account classes needed for structured recording, not a complete SME accounting interface. Only expense-management workflows are committed. A local label is separate from its category identity and account mapping.')
 r.table('Illustrative local category mappings',['User-facing category','Internal accounting account'],[(x,y) for x,y in [('Ration','Groceries Expense'),('Transport','Transport Expense'),('Dining','Dining Expense'),('Utilities','Utilities Expense'),('Healthcare','Healthcare Expense'),('Education','Education Expense'),('Entertainment','Entertainment Expense'),('Household Supplies','Household Supplies Expense')]],[2.4,3.62])
 r.p('These mappings illustrate the approved architecture; account codes and the final hierarchy are later design decisions. For groceries purchased in cash for Rs. 2,500, a double-entry implementation debits Groceries Expense by Rs. 2,500 and credits Cash by Rs. 2,500. Every journal must balance wherever double-entry is used. Users should not enter journals or account codes. No accounting or regulatory certification is claimed.')
 r.p('A mixed receipt retains its receipt total and separately categorised line items. Category totals must not count the same purchase once at transaction level and again at line-item level. Edits and deletion must also update dependent totals and price history consistently. Accounting corrections must preserve balance where journals are used. These are project-team design constraints derived from the need for trustworthy records.')
 r.h('Semantic product classification',2)
 r.p('Original receipt text, individual product identity, brand, canonical product type, expense category and accounting classification are separate concepts. Preserve the original text alongside corrected values and inferred classifications. A canonical product type groups related products without replacing their identities.')
 r.table('Illustrative semantic product distinctions',['Original text and brand','Product identity and type','Category and account'],[('Surf Excel 1kg\nBrand: Surf Excel','Distinct Surf Excel 1kg product\nLaundry Detergent','Household Supplies\nHousehold Supplies Expense'),('Bonus 1kg\nBrand: Bonus','Distinct Bonus 1kg product\nLaundry Detergent','Household Supplies\nHousehold Supplies Expense')],[1.65,2.3,2.07])
 r.p('Both products can support analysis of Laundry Detergent spending. Exact price history still distinguishes brand, product and pack size. Comparisons require comparable quantities and units; where these are missing, show recorded prices with a limitation instead of claiming equivalent value. Uncertain product matches remain reviewable. Semantic similarity is not proof that two products are identical.')
 r.h('AI Financial Assistant',2)
 r.p('The AI Financial Assistant is a core conversational capability grounded in the current user\'s authorised DailyKhata records. It supports natural-language questions, spending summaries, comparisons between periods and categories, recurring-spending analysis, product-price comparisons, spending patterns, unusually high recorded spending, explanations of change and possible saving opportunities. It may give personalised general financial suggestions based on available historical data.')
 r.p('Illustrative response, not research: restaurant spending of Rs. 8,400 this month compared with Rs. 5,900 last month gives a difference of Rs. 2,500. The assistant may explain that returning to the earlier level would reduce spending by that amount. It must distinguish recorded facts, calculations and suggestions, state the relevant period, and acknowledge missing data. It does not move money, execute payments or control accounts.')
 r.p('Conversational analysis of repeated past expenses is core; automatically generating recurring expenses is future scope. Data-grounded observations about high spending are core; predictive anomaly detection and advanced financial recommendation systems remain future scope. The team may integrate existing models and services instead of training a foundation model.')
 r.h('Localisation and delivery limits',2)
 r.p('English and Urdu localisation is Should Have. The architecture should permit translated labels, right-to-left Urdu presentation, local terminology, Urdu-friendly interaction and potential Urdu interaction with the AI Financial Assistant. An initial English interface does not prove bilingual localisation. Urdu support is not guaranteed until implemented and validated. Local category labels remain Must Have independently of translation.')
 r.p('The application is mobile-first and should support ordinary Android-class devices. Receipt OCR targets reasonably clear printed receipts in supported formats, with correction and manual fallback for failures. No universal receipt accuracy or response-time threshold has been validated. Frameworks, model providers, device test specifications and measurable performance targets are later decisions. Offline draft entry is Could Have and does not imply offline OCR or offline assistant services.')

def boundaries(r):
 r.p('Primary users are homemakers and household budget managers, university students and salaried professionals. Small shopkeepers are secondary direct users for petty cash, supplier expenses, receipt capture and expense summaries. Family members and advisers are secondary beneficiaries of information shared by the account holder.')
 r.p('The baseline is a single-account expense-management system. Household expenses can be recorded by the account holder and reviewed together. CSV and/or PDF export is Should Have; a dedicated exportable household summary is Could Have. Shared household accounts, invitations, household permissions and collaborative editing are not current requirements.')
 r.h('Current exclusions',2); r.p(EXCLUDED)
 r.h('Future scope',2); r.p(FUTURE)
 r.p('A feature can be excluded from the current project and also be named as potential future work. Future listing is not a delivery promise. Recording an expense paid by a method described by the user does not create a bank or wallet integration.')

def references(r,full=False):
 r.h('References')
 refs=[('[S1]','DailyKhata. Target Market & Audience Analysis. Original submitted D1, dailykhata-d1.docx and corresponding six-page PDF, 01_Deliverable_1. Undated.'),('[S2]','DailyKhata. Original submitted D2, dailykhata-d2.pdf, 02_Deliverable_2, 12 pages. Cover repeats the D1 title; body contains problem identification, stakeholder identification and requirement elicitation. Undated.'),('[A1]','Deliverable 1 task.txt, 01_Deliverable_1/task. User, problem, expectation and target-market assessment criteria. Undated.'),('[A2]','Deliverable 2 task.txt, 02_Deliverable_2/task. Official D2 title and evidence-based assignment requirements. Stated deadline: 25 August 2026, not proof of submission.'),('[G1]','DailyKhata Scope Baseline v1.0 and DailyKhata Documentation Standard v1.0, 00_Project_Governance. Both dated 23 September 2026.'),('[U1]','Project owner instructions for reconstruction, 23 September 2026. Authoritative scope, team identifiers, priorities and constraints.')]
 if full: refs += [('[A3]','Deliverable 3 Requirements Specification, 03_Deliverable_3/task/Deliverable 3.pdf, five pages. Read for compatibility only.'),('[M1]','Design Thinking.pdf, 02_Deliverable_2/task, six instructional slides. Observe, interview, distinguish observation from interpretation, and identify needs.'),('[M2]','Important Links.txt, 02_Deliverable_2/task, 19 teaching-video links. Retained as teaching references; video contents are not treated as research evidence.')]
 for key,t in refs: r.p(key+' '+t)

def d1():
 r=Report('Target Market & Audience Analysis','DailyKhata_Deliverable_1_v2.0.docx','01_Deliverable_1',1)
 r.h('Purpose and project foundation'); r.p('DailyKhata addresses the effort of maintaining useful personal and household expense records. This deliverable establishes its intended audience, user problems, needs, expectations and target-market justification. It is the factual and scope foundation inherited by Deliverable 2. It does not claim an implemented system or completed user validation. [A1, U1]'); r.p(DEFINITION)
 r.h('Market and Pakistani context'); r.p('The industry is financial technology, focused on personal expense management. The relevant market is expense-management applications for Pakistani users. The selected segment comprises adults who manage personal or household spending, use smartphones, and regularly encounter paper receipts or cash purchases without receipts. The original documents use an indicative age range of 18 to 50; this is a design focus, not a verified market statistic or a maximum user age. [S1, S2]')
 r.p('A grocery receipt can contain several products that later need different expense categories. A rickshaw fare may leave no receipt at all. A notebook or phone note can preserve an amount, but the user still has to maintain consistent entries and assemble totals. DailyKhata combines capture and review with structured records so that the same information supports later search, summaries and questions. This is project-team analysis of the reported workflows, not a claim that every Pakistani household uses the same method.')
 r.h('Evidence and market claim limits',2); r.p('The strongest available project evidence is the original D2 narrative: three reported interviews, informal observation and a reported comparison of existing applications. It describes effort, delayed entry and distrust of incorrect categories. No interview transcripts, observation counts, recordings, survey data or competitor test artefacts accompany that narrative. These findings support a small qualitative design rationale, not prevalence or market-size estimates. D2 retains the detailed evidence register. [S2, pp. 3, 10-12]')
 r.p('The original estimates of several million potential users and typical monthly spending of Rs. 40,000 are not retained as facts because no supporting calculation or source is present. Statements that all competing applications require manual entry or lack price history are also not adopted. The market opportunity is framed as a testable fit between the selected users and DailyKhata\'s workflow, not a verified claim of market uniqueness. [DEC-08, DEC-09]')
 r.h('Primary and secondary users'); boundaries(r)
 r.h('Canonical personas'); persona_section(r)
 r.h('User problems needs and expectations')
 r.p(PROBLEM)
 r.table('User problems linked to needs and benefits',['Problem and source','Need and response','Expected benefit'],[
 ('Scattered or missing receipt records [S1; S2 p. 3]','Receipt digitisation, reviewed OCR and searchable history','Less transcription and a retrievable structured record'),
 ('Small no-receipt cash purchases omitted [S2 pp. 10-11]','Manual quick entry and text-based cash entry','Record spending at purchase time or afterwards'),
 ('Manual classification and incorrect categories [S2 p. 10]','Automatic categorisation with confirmation and correction','Reduce repeated work while retaining user control'),
 ('Difficulty reconstructing totals [S2 pp. 3, 10]','Category totals, monthly summaries and dashboard','Understand recorded spending without rereading every note'),
 ('No usable item comparison in reported workflow [S1; S2 p. 10]','Item-level price history and semantic product classification','Compare recorded products without losing brand identity'),
 ('Effort finding an answer [S1]','AI Financial Assistant grounded in authorised records','Ask plain-language questions about spending'),
 ('Sensitive financial records [S1; U1]','Authentication, access controls and secure storage','Keep each user\'s data private')],[1.85,2.25,1.92])
 r.p('Users expect low effort, understandable values and reliable correction. Automation should reduce transcription rather than hide uncertainty. The record is saved after review, and saved transactions remain editable and deletable. A failed OCR attempt should lead to a clear retry or manual-entry path. Receipt image archiving is optional, so structured record retention must not be described as permanent preservation of every receipt image.')
 r.p('Fast interaction is an expectation, not a measured result. The source interview question about logging in under 10 seconds establishes interest in low-effort capture, not a validated end-to-end benchmark. Later evaluation should separate simple cash entry from image processing and user correction. Privacy, ordinary-device usability and coherent totals are essential to trust. [DEC-10, DEC-11]')
 r.h('Authoritative capability baseline'); scope_table(r)
 r.h('Foundational design concepts'); architecture(r)
 r.h('Target market justification and feasibility'); r.p('The selected groups face related recording tasks at different scales. Ayesha needs weekly household visibility, Hamza needs a short path for frequent cash purchases, and Bilal needs a coherent monthly view. One expense ledger, supported by receipt and text capture, can serve these needs without separate enterprise workflows. Usman can reuse that same ledger for expenses without changing the primary market. [S1, S2; project-team analysis]')
 r.p('A four-student team can divide work into capture and review, ledger and accounting mappings, product semantics and price history, and assistant and presentation, while sharing integration and testing. These are proposed work areas, not assignments to named students. Using existing OCR and language-model components is permitted. The academic contribution lies in integration, structured data, correction, semantic modelling, traceability and evaluation. Optional features should follow a working end-to-end Must Have workflow.')
 r.p('The fit should be tested with representative users and receipts. Useful validation includes whether users can correct extraction, record a no-receipt expense, find a past transaction and understand an assistant answer. Record actual task completion and error causes before setting quantitative acceptance targets. No market demand, retention improvement or OCR accuracy has yet been established by this reconstruction.')
 r.h('Conclusion'); r.p('DailyKhata has a focused personal and household expense-management baseline. Its receipt and cash workflows connect to structured accounting records, item-level price history and a data-grounded AI Financial Assistant. D2 inherits the same personas, scope identifiers and priorities, and expands the evidence and elicitation rationale. Research verification and supervisor review remain distinct from document consistency.')
 references(r); r.finish()

if __name__=='__main__':
 if sys.argv[-1]=='d1': d1()
